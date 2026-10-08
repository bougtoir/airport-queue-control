from __future__ import annotations

import csv
import shutil
from collections import Counter
from dataclasses import replace
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from pptx import Presentation
from pptx.util import Inches, Pt

from airport_batch_queue.config import load_config
from airport_batch_queue.simulator import run_simulation

COLORS = {
    "S0": "#0072B2",
    "S3": "#D55E00",
    "S5": "#009E73",
    "neutral": "#666666",
}
FIGURE_DIR = Path("figures")
SUPPLEMENT_DIR = Path("supplementary_figures")
TABLE_DIR = Path("tables")
SUPPLEMENT_TABLE_DIR = Path("supplementary_tables")
DATA_DIR = Path("results/phase10")


def _prepare() -> None:
    for path in (
        FIGURE_DIR,
        SUPPLEMENT_DIR,
        TABLE_DIR,
        SUPPLEMENT_TABLE_DIR,
        DATA_DIR,
    ):
        path.mkdir(parents=True, exist_ok=True)


def _save_figure(figure: plt.Figure, name: str, supplement: bool = False) -> Path:
    directory = SUPPLEMENT_DIR if supplement else FIGURE_DIR
    png_path = directory / f"{name}.png"
    svg_path = directory / f"{name}.svg"
    pdf_path = directory / f"{name}.pdf"
    figure.savefig(png_path, dpi=300, bbox_inches="tight")
    figure.savefig(svg_path, bbox_inches="tight")
    figure.savefig(pdf_path, bbox_inches="tight")
    svg_path.write_text(
        "\n".join(line.rstrip() for line in svg_path.read_text(encoding="utf-8").splitlines())
        + "\n",
        encoding="utf-8",
    )
    plt.close(figure)
    return png_path


def _concept_figure() -> tuple[Path, str]:
    figure, axes = plt.subplots(1, 3, figsize=(12, 3.5))
    titles = ["S0: serpentine", "S3: stationary rows", "Threshold release"]
    for axis, title in zip(axes, titles, strict=True):
        axis.set_title(title, fontweight="bold")
        axis.set_xlim(0, 10)
        axis.set_ylim(0, 7)
        axis.axis("off")
    axes[0].plot([1, 9, 1, 9, 1, 9], [1, 1, 2, 3, 4, 5], color=COLORS["S0"], lw=3)
    for x, y in zip([2, 4, 6, 8], [1, 2, 3, 4], strict=True):
        axes[0].add_patch(plt.Circle((x, y), 0.15, color=COLORS["S0"]))
    for y in (1.5, 3, 4.5):
        axes[1].plot([1, 8], [y, y], color=COLORS["neutral"], lw=2)
        for x in np.linspace(2, 7, 4):
            axes[1].add_patch(plt.Circle((x, y), 0.15, color=COLORS["S3"]))
    axes[1].add_patch(
        FancyArrowPatch((8, 3), (9.5, 3), arrowstyle="->", mutation_scale=16)
    )
    for index, label in enumerate(("Wait", "Sense", "Release", "Screen")):
        x = 0.4 + index * 2.4
        axes[2].add_patch(
            FancyBboxPatch(
                (x, 2.6),
                1.8,
                1.2,
                boxstyle="round,pad=0.1",
                facecolor="#E6E6E6",
                edgecolor=COLORS["neutral"],
            )
        )
        axes[2].text(x + 0.9, 3.2, label, ha="center", va="center")
        if index < 3:
            axes[2].add_patch(
                FancyArrowPatch(
                    (x + 1.8, 3.2),
                    (x + 2.35, 3.2),
                    arrowstyle="->",
                    mutation_scale=14,
                )
            )
    figure.suptitle("Separating waiting from repeated movement")
    caption = "Conceptual control architecture; diagrams are generic and not site layouts."
    return _save_figure(figure, "figure_01_concept_architecture"), caption


def _trajectory_figure() -> tuple[Path, str]:
    base = load_config("configs/default.yaml")
    results = {}
    for policy, layout in (("S0", "B0"), ("S3", "B1")):
        config = replace(
            base,
            horizon_seconds=3600,
            drain_limit_seconds=3600,
            control=replace(base.control, policy=policy),
            geometry=replace(base.geometry, layout=layout),
        )
        results[policy] = run_simulation(config)
    figure, axes = plt.subplots(1, 2, figsize=(11, 4))
    for policy in ("S0", "S3"):
        trajectory = pd.DataFrame(results[policy].trajectories)
        axes[0].plot(
            trajectory["time_seconds"] / 60,
            trajectory["upstream_waiting"],
            label=policy,
            color=COLORS[policy],
        )
    axes[0].set(
        xlabel="Time (min)",
        ylabel="Upstream waiting passengers",
        title="Representative queue trajectories",
    )
    axes[0].legend(frameon=False)
    metrics = ["Movement starts", "Movement stops", "Distance (m)", "Turning / 100°"]
    values = {}
    for policy in ("S0", "S3"):
        summary = results[policy].summary
        values[policy] = [
            float(summary["mean_movement_starts"]),
            float(summary["mean_movement_stops"]),
            float(summary["mean_movement_distance_m"]),
            float(summary["mean_turning_angle_degrees"]) / 100,
        ]
    x = np.arange(len(metrics))
    axes[1].bar(x - 0.18, values["S0"], 0.36, label="S0", color=COLORS["S0"])
    axes[1].bar(x + 0.18, values["S3"], 0.36, label="S3", color=COLORS["S3"])
    axes[1].set_xticks(x, metrics, rotation=20, ha="right")
    axes[1].set_title("Movement-event decomposition")
    axes[1].legend(frameon=False)
    figure.tight_layout()
    caption = "One common-seed realization; movement endpoints are modeled locomotion metrics."
    return _save_figure(figure, "figure_02_trajectories_events"), caption


def _movement_load_figures() -> tuple[tuple[Path, str], tuple[Path, str]]:
    data = pd.read_csv("results/phase5/replications.csv")
    data = data[data["policy"].isin(["S0", "S3"])]
    endpoints = [
        ("mean_movement_starts", "Movement starts"),
        ("mean_movement_stops", "Movement stops"),
        ("mean_movement_distance_m", "Distance (m)"),
        ("mean_turning_angle_degrees", "Turning (degrees)"),
    ]
    output_rows = []
    summaries: dict[str, dict[str, tuple[pd.Series, pd.Series]]] = {}
    for column, _ in endpoints:
        summaries[column] = {}
        for policy in ("S0", "S3"):
            grouped = data[data["policy"] == policy].groupby("load_regime")[column]
            means = grouped.mean()
            errors = 1.96 * grouped.sem()
            summaries[column][policy] = (means, errors)
            output_rows.extend(
                {
                    "endpoint": column,
                    "policy": policy,
                    "load_regime": load,
                    "mean": value,
                    "ci95_half_width": error,
                }
                for load, value, error in zip(
                    means.index,
                    means.values,
                    errors.values,
                    strict=True,
                )
            )
    pd.DataFrame(output_rows).to_csv(DATA_DIR / "figure_03_data.csv", index=False)

    styles = {
        "S0": {"marker": "o", "linestyle": "-"},
        "S3": {"marker": "s", "linestyle": "--"},
    }
    figure, axis = plt.subplots(figsize=(7.2, 4.4))
    for policy in ("S0", "S3"):
        means, errors = summaries["mean_movement_starts"][policy]
        axis.errorbar(
            means.index,
            means.values,
            yerr=errors.values,
            label=policy,
            color=COLORS[policy],
            linewidth=2,
            capsize=3,
            **styles[policy],
        )
    axis.set(
        xlabel="Nominal load ratio",
        ylabel="Modeled movement starts per passenger",
        title="Model-dependent movement starts across nominal load",
    )
    axis.grid(alpha=0.2)
    axis.legend(frameon=False)
    figure.tight_layout()
    main_caption = (
        "Model-dependent mean movement starts and 95% Monte Carlo intervals; "
        "S3 has one start by construction under the specified progression rule."
    )
    main_item = (
        _save_figure(figure, "figure_03_movement_by_load"),
        main_caption,
    )

    figure, axes = plt.subplots(2, 2, figsize=(10, 7), sharex=True)
    classifications = (
        "Model-dependent dynamic",
        "Model-dependent dynamic",
        "Structural route outcome",
        "Structural route outcome",
    )
    for axis, (column, label), classification in zip(
        axes.flat,
        endpoints,
        classifications,
        strict=True,
    ):
        for policy in ("S0", "S3"):
            means, errors = summaries[column][policy]
            axis.errorbar(
                means.index,
                means.values,
                yerr=errors.values,
                label=policy,
                color=COLORS[policy],
                linewidth=2,
                capsize=3,
                **styles[policy],
            )
        axis.set_ylabel(label)
        axis.set_title(classification, fontsize=10)
        axis.grid(alpha=0.2)
    for axis in axes[-1]:
        axis.set_xlabel("Nominal load ratio")
    axes[0, 0].legend(frameon=False)
    figure.suptitle("Complete dynamic and structural movement endpoints")
    figure.tight_layout()
    supplement_caption = (
        "Complete movement endpoint panel. Starts and stops are model-dependent dynamic "
        "outcomes; distance and turning are structural route outcomes. Means and 95% Monte "
        "Carlo intervals are shown."
    )
    supplement_item = (
        _save_figure(
            figure,
            "supplement_07_full_movement_endpoints",
            supplement=True,
        ),
        supplement_caption,
    )
    return main_item, supplement_item


def _control_response_figure() -> tuple[Path, str]:
    data = pd.read_csv("results/phase7/replications.csv")
    data = data[
        (data["policy"] == "S3")
        & (data["layout"] == "B1")
        & (data["row_angle_degrees"] == 0)
        & (data["load_regime"] == 0.95)
    ]
    data = data.groupby(
        ["batch_size", "downstream_threshold"],
        as_index=False,
    ).agg(
        throughput_per_hour=("throughput_per_hour", "mean"),
        utilization=("utilization", "mean"),
        starvation_fraction_with_upstream_demand=(
            "starvation_fraction_with_upstream_demand",
            "mean",
        ),
    )
    metrics = [
        ("throughput_per_hour", "Throughput (passengers/h)"),
        ("utilization", "Lane utilization"),
        (
            "starvation_fraction_with_upstream_demand",
            "Starvation fraction",
        ),
    ]
    figure, axes = plt.subplots(1, 3, figsize=(12, 3.7))
    for axis, (column, label) in zip(axes, metrics, strict=True):
        for threshold, values in data.groupby("downstream_threshold"):
            values = values.sort_values("batch_size")
            axis.plot(
                values["batch_size"],
                values[column],
                marker="o",
                label=f"Threshold {threshold}",
            )
        axis.set(xlabel="Batch size", ylabel=label)
        axis.grid(alpha=0.2)
    axes[0].legend(frameon=False, fontsize=8)
    figure.suptitle("Controller settings change operational performance")
    figure.tight_layout()
    data.to_csv(DATA_DIR / "figure_04_data.csv", index=False)
    caption = "S3/B1 designs at load 0.95; endpoints remain separate rather than weighted."
    return _save_figure(figure, "figure_04_control_response"), caption


def _space_figure() -> tuple[Path, str]:
    data = pd.read_csv("results/phase3_space_guidance/layout_angle_screen.csv")
    data = (
        data.groupby(["layout", "row_angle_degrees"], as_index=False)
        .agg(
            required_area=(
                "required_waiting_area_m2_at_density_limit",
                "mean",
            ),
            capacity=("acceptable_waiting_capacity", "mean"),
        )
    )
    figure, axes = plt.subplots(1, 2, figsize=(10, 4))
    for layout, values in data.groupby("layout"):
        axes[0].plot(
            values["row_angle_degrees"],
            values["required_area"],
            marker="o",
            label=layout,
        )
        axes[1].plot(
            values["row_angle_degrees"],
            values["capacity"],
            marker="o",
            label=layout,
        )
    axes[0].set(xlabel="Row angle (degrees)", ylabel="Required waiting area (m²)")
    axes[1].set(xlabel="Row angle (degrees)", ylabel="Acceptable waiting capacity")
    axes[0].legend(frameon=False, fontsize=8)
    for axis in axes:
        axis.grid(alpha=0.2)
    figure.suptitle("Space requirements depend on geometry and orientation")
    figure.tight_layout()
    data.to_csv(DATA_DIR / "figure_05_data.csv", index=False)
    caption = "Phase 3 generic geometry screen averaged across seeds and footprint cases."
    return _save_figure(figure, "figure_05_space_geometry"), caption


def _pareto_figure() -> tuple[Path, str]:
    all_designs = pd.read_csv("results/phase7/design_summary.csv")
    frontier = pd.read_csv("results/phase7/pareto_front.csv")
    throughput_min = float(all_designs["throughput_per_hour"].min())
    throughput_range = float(all_designs["throughput_per_hour"].max() - throughput_min)

    def marker_size(values: pd.Series) -> pd.Series:
        return 18 + 72 * (values - throughput_min) / throughput_range

    figure, axes = plt.subplots(
        1,
        2,
        figsize=(12, 5.2),
        gridspec_kw={"width_ratios": [1.35, 1]},
    )
    evaluated = axes[0].scatter(
        all_designs["mean_wait_seconds"],
        all_designs["required_waiting_area_m2_at_density_limit"],
        c=all_designs["mean_movement_starts"],
        s=marker_size(all_designs["throughput_per_hour"]),
        cmap="cividis",
        alpha=0.55,
        marker="o",
        edgecolor="#666666",
        linewidth=0.35,
        label="All evaluated designs",
    )
    axes[0].scatter(
        frontier["mean_wait_seconds"],
        frontier["required_waiting_area_m2_at_density_limit"],
        c=frontier["mean_movement_starts"],
        s=marker_size(frontier["throughput_per_hour"]) + 55,
        cmap=evaluated.cmap,
        norm=evaluated.norm,
        marker="D",
        edgecolor="black",
        linewidth=1.2,
        label="Nondominated designs",
    )
    axes[0].set(
        xlabel="Mean wait (s)",
        ylabel="Required waiting area (m²)",
        title="(a) Evaluated objective space",
    )
    axes[0].legend(frameon=False, loc="upper left")
    axes[0].grid(alpha=0.2)
    figure.colorbar(
        evaluated,
        ax=axes[0],
        label="Modeled movement starts per passenger",
    )

    objective_columns = [
        "mean_wait_seconds",
        "required_waiting_area_m2_at_density_limit",
        "mean_movement_starts",
        "throughput_per_hour",
    ]
    minima = all_designs[objective_columns].min()
    spans = all_designs[objective_columns].max() - minima
    normalized = (frontier[objective_columns] - minima) / spans
    normalized["throughput_per_hour"] = 1 - normalized["throughput_per_hour"]
    image = axes[1].imshow(
        normalized.to_numpy(),
        cmap="cividis",
        vmin=0,
        vmax=1,
        aspect="auto",
    )
    axes[1].set_xticks(
        range(len(objective_columns)),
        ["Wait", "Area", "Starts", "Throughput\nshortfall"],
    )
    axes[1].set_yticks(
        range(len(frontier)),
        [f"P{index}" for index in range(1, len(frontier) + 1)],
    )
    axes[1].set_title("(b) Nondominated objective profiles")
    actual_formats = (".0f", ".1f", ".2f", ".0f")
    for row_index, row in enumerate(
        frontier[objective_columns].itertuples(index=False, name=None)
    ):
        for column_index, (value, value_format) in enumerate(
            zip(row, actual_formats, strict=True)
        ):
            axes[1].text(
                column_index,
                row_index,
                format(value, value_format),
                ha="center",
                va="center",
                color="white" if normalized.iloc[row_index, column_index] > 0.55 else "black",
                fontsize=8,
            )
    figure.colorbar(
        image,
        ax=axes[1],
        label="Normalized burden within evaluated range",
    )
    figure.suptitle("Movement-aware Pareto trade-offs")
    figure.tight_layout()
    figure_data = all_designs[
        [
            "design_id",
            "mean_wait_seconds",
            "required_waiting_area_m2_at_density_limit",
            "mean_movement_starts",
            "throughput_per_hour",
            "pareto_efficient",
        ]
    ].copy()
    figure_data.to_csv(DATA_DIR / "figure_06_data.csv", index=False)
    caption = (
        "All evaluated designs are circles: x is mean wait, y is conditional required "
        "waiting area, monotonic-lightness color is modeled starts/passenger, and marker "
        "area is throughput. The eight nondominated designs are black-edged diamonds and "
        "their normalized objective profiles are shown at right; all share one modeled "
        "start/passenger under the current progression rule."
    )
    return _save_figure(figure, "figure_06_pareto_front"), caption


def _phase_map_figure() -> tuple[Path, str]:
    data = pd.read_csv("results/phase7/design_summary.csv")
    data = data[data["policy"] == "S3"].copy()
    baseline_starts = float(
        pd.read_csv("results/phase7/design_summary.csv")
        .query("policy == 'S0'")["mean_movement_starts"]
        .iloc[0]
    )
    data = data[data["mean_movement_starts"] < baseline_starts]
    phase = data.pivot_table(
        index="downstream_threshold",
        columns="batch_size",
        values="minimum_throughput_relative_lower_bound",
        aggfunc="min",
    )
    figure, axis = plt.subplots(figsize=(6, 4.5))
    image = axis.imshow(phase.values, cmap="coolwarm", aspect="auto")
    axis.set_xticks(range(len(phase.columns)), [str(value) for value in phase.columns])
    axis.set_yticks(range(len(phase.index)), [str(value) for value in phase.index])
    axis.set(xlabel="Batch size", ylabel="Downstream threshold")
    axis.set_title("Worst throughput bound among movement-improving variants")
    for row in range(phase.shape[0]):
        for column in range(phase.shape[1]):
            axis.text(
                column,
                row,
                f"{phase.iloc[row, column]:.2f}",
                ha="center",
                va="center",
                color="black",
            )
    figure.colorbar(image, ax=axis, label="Paired relative lower bound")
    figure.tight_layout()
    phase.to_csv(DATA_DIR / "figure_07_data.csv")
    caption = "All cells improve movement starts; the acceptability threshold is −0.02."
    return _save_figure(figure, "figure_07_phase_map"), caption


def _supplementary_figures() -> list[tuple[Path, str]]:
    outputs: list[tuple[Path, str]] = []
    stress = pd.read_csv("results/phase8/scenario_comparisons.csv")
    selections = [
        (
            "supplement_01_arrival_service_sensitivity",
            stress[stress["category"].isin(["arrival_process", "service_variability"])],
            "Arrival-process and service-variability sensitivity",
        ),
        (
            "supplement_02_compliance_gate_delay",
            stress[stress["category"].isin(["compliance", "guidance_delay"])],
            "Compliance and guidance-delay sensitivity",
        ),
    ]
    for name, values, title in selections:
        figure, axis = plt.subplots(figsize=(8, 4))
        axis.bar(
            values["scenario_id"],
            values["throughput_relative_lower_bound"],
            color="#56B4E9",
        )
        axis.axhline(-0.02, color=COLORS["S3"], ls="--", label="−2% criterion")
        axis.set(ylabel="Paired throughput lower bound", title=title)
        axis.tick_params(axis="x", rotation=35)
        axis.legend(frameon=False)
        figure.tight_layout()
        outputs.append(
            (
                _save_figure(figure, name, supplement=True),
                "Phase 8 paired throughput lower bounds.",
            )
        )
    geometry = pd.read_csv("results/phase3_space_guidance/layout_angle_screen.csv")
    for name, column, title, ylabel in (
        (
            "supplement_03_footprint_accessibility",
            "required_waiting_area_m2_at_density_limit",
            "Footprint sensitivity",
            "Required waiting area (m²)",
        ),
        (
            "supplement_04_row_angle",
            "mean_movement_distance_m",
            "Row-angle sensitivity",
            "Mean movement distance (m)",
        ),
    ):
        figure, axis = plt.subplots(figsize=(8, 4))
        grouped = geometry.groupby(
            ["footprint_case" if "footprint" in name else "layout", "row_angle_degrees"]
        )[column].mean()
        for label, values in grouped.groupby(level=0):
            series = values.droplevel(0)
            axis.plot(series.index, series.values, marker="o", label=label)
        axis.set(xlabel="Row angle (degrees)", ylabel=ylabel, title=title)
        axis.legend(frameon=False)
        figure.tight_layout()
        outputs.append(
            (
                _save_figure(figure, name, supplement=True),
                "Generic geometry sensitivity; accessibility minima remain constraints.",
            )
        )
    failures = stress.copy()
    flag_columns = [
        "throughput_failure",
        "movement_null",
        "waiting_failure",
        "space_failure",
        "release_surge_failure",
        "conflict_proxy_failure",
        "staffing_failure",
    ]
    matrix = (
        failures.groupby("category")[flag_columns]
        .agg(lambda series: (series == "yes").mean())
        .sort_index()
    )
    figure, axis = plt.subplots(figsize=(9, 5))
    image = axis.imshow(matrix.values, vmin=0, vmax=1, cmap="magma", aspect="auto")
    axis.set_xticks(
        range(len(flag_columns)),
        [value.replace("_failure", "").replace("_", " ") for value in flag_columns],
        rotation=35,
        ha="right",
    )
    axis.set_yticks(range(len(matrix.index)), matrix.index)
    axis.set_title("Failure criteria by stress-test category")
    figure.colorbar(image, ax=axis, label="Fraction triggering")
    figure.tight_layout()
    outputs.append(
        (
            _save_figure(
                figure,
                "supplement_05_falsification_map",
                supplement=True,
            ),
            "Proxy criteria are engineering screens, not validated safety thresholds.",
        )
    )
    transfer = pd.read_csv("results/phase9/transferability_comparisons.csv")
    figure, axis = plt.subplots(figsize=(9, 4))
    positions = np.arange(len(transfer["scenario_id"].unique()))
    for offset, policy in ((-0.18, "S3"), (0.18, "S5")):
        values = transfer[transfer["policy"] == policy]
        axis.bar(
            positions + offset,
            values["throughput_relative_lower_bound"],
            width=0.36,
            label=policy,
            color=COLORS[policy],
        )
    axis.axhline(-0.02, color="black", ls="--")
    axis.set_xticks(
        positions,
        transfer["scenario_id"].unique(),
        rotation=35,
        ha="right",
    )
    axis.set(ylabel="Paired throughput lower bound", title="Generic-geometry benchmark")
    axis.legend(frameon=False)
    figure.tight_layout()
    outputs.append(
        (
            _save_figure(
                figure,
                "supplement_06_multiple_geometries",
                supplement=True,
            ),
            "Finite benchmark results do not establish universal transferability.",
        )
    )
    return outputs


def _table_document(
    title: str,
    headers: list[str],
    rows: list[list[str]],
    path: Path,
) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value
    document.save(path)


def _multi_panel_table_document(
    title: str,
    data: list[dict[str, str]],
    panels: list[tuple[str, list[tuple[str, str]]]],
    path: Path,
) -> None:
    document = Document()
    section = document.sections[0]
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    document.add_heading(title, level=1)
    for panel_title, columns in panels:
        document.add_heading(panel_title, level=2)
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for cell, (_, label) in zip(table.rows[0].cells, columns, strict=True):
            cell.text = label
        for row in data:
            cells = table.add_row().cells
            for cell, (column, _) in zip(cells, columns, strict=True):
                cell.text = row[column]
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
    document.save(path)


def _display_number(value: float, digits: int) -> str:
    rendered = f"{value:.{digits}f}"
    return rendered.replace("-", "−")


def _build_tables() -> None:
    for stale in TABLE_DIR.glob("table_01_parameters_sources.*"):
        stale.unlink()
    parameters = pd.read_csv("references/parameter_sources.csv")
    base = load_config("configs/default.yaml")

    def parameter_row(name: str) -> pd.Series:
        return parameters.loc[parameters["parameter"] == name].iloc[0]

    walking = parameter_row("airport_free_walking_speed")
    startup = parameter_row("pedestrian_startup_delay")
    reaction = parameter_row("pedestrian_reaction_delay")
    slowdown = parameter_row("pedestrian_slowdown_time")
    service = parameter_row("screening_service_rate")
    route_width = parameter_row("accessible_route_clear_width")
    turning = parameter_row("accessible_turning_diameter")
    luggage = parameter_row("luggage_footprint_per_bag")
    summary_rows = [
        [
            "Demand and screening",
            (
                f"{base.arrival.rate_per_hour:.0f} passengers/h; "
                f"{base.service.servers} lanes; {base.service.mean_seconds:.0f} s mean service"
            ),
            f"{service.lower:.0f}–{service.upper:.0f} passengers/h/lane",
            "Synthetic demand; guidance-informed engineering range",
        ],
        [
            "Walking speed",
            f"{base.passenger.walking_speed_mps:.2f} m/s",
            f"{walking.lower:.2f}–{walking.upper:.2f} m/s",
            "Peer-reviewed airport observations",
        ],
        [
            "Stop/start timing",
            (
                f"{startup.estimate:.2f} s startup; {reaction.estimate:.2f} s reaction; "
                f"{slowdown.estimate:.2f} s slowdown"
            ),
            (
                f"{startup.lower:.2f}–{startup.upper:.2f}; "
                f"{reaction.lower:.2f}–{reaction.upper:.2f}; "
                f"{slowdown.lower:.2f}–{slowdown.upper:.2f} s"
            ),
            "Peer-reviewed controlled motion capture",
        ],
        [
            "Passenger heterogeneity",
            (
                f"{base.passenger.luggage_prevalence:.2f} luggage prevalence; "
                f"{base.passenger.reduced_mobility_prevalence:.2f} reduced mobility"
            ),
            "0–1.00; 0–0.20 proportions",
            "Explicit assumptions; stress tested",
        ],
        [
            "Luggage space",
            f"{base.passenger.luggage_footprint_m2_per_bag:.2f} m²/bag",
            f"{luggage.lower:.2f}–{luggage.upper:.2f} m²/bag",
            "Literature-informed engineering assumption",
        ],
        [
            "S3 release rule",
            (
                f"batch {base.control.batch_size}; threshold "
                f"{base.control.downstream_threshold}; maximum hold "
                f"{base.control.maximum_hold_seconds:.0f} s"
            ),
            "Prespecified Phase 7 control grid",
            "Explicit controller design factors",
        ],
        [
            "Configured routes",
            (
                f"S0 {base.geometry.serpentine_distance_m:.0f} m/"
                f"{base.geometry.serpentine_turns} turns; "
                f"stationary {base.geometry.stationary_distance_m:.0f} m/"
                f"{base.geometry.stationary_turns} turns"
            ),
            "Matched short, mid, and long routes",
            "Structural assumptions; matched-route sensitivity",
        ],
        [
            "Spatial constraints",
            (
                f"{base.geometry.accessible_route_width_m:.2f} m route; "
                f"{base.geometry.turning_diameter_m:.2f} m turning; "
                f"{base.geometry.max_waiting_density_per_m2:.2f} passengers/m²"
            ),
            (
                f"minimum {route_width.estimate:.3f} m route and "
                f"{turning.estimate:.3f} m turning"
            ),
            "Official accessibility minima plus engineering assumptions",
        ],
    ]
    summary_headers = ["Parameter/group", "Base value", "Range", "Evidence/type"]
    with (TABLE_DIR / "table_01_parameter_summary.csv").open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(summary_headers)
        writer.writerows(summary_rows)
    _table_document(
        "Table 1. Reader-facing model parameter summary",
        summary_headers,
        summary_rows,
        TABLE_DIR / "table_01_parameter_summary.docx",
    )
    parameters.to_csv(
        SUPPLEMENT_TABLE_DIR / "table_s1_full_parameter_provenance.csv",
        index=False,
    )
    parameter_records = [
        {key: "" if pd.isna(value) else str(value) for key, value in row.items()}
        for row in parameters.to_dict(orient="records")
    ]
    _multi_panel_table_document(
        "Table S1. Complete parameter provenance",
        parameter_records,
        [
            (
                "Table S1A. Values, ranges, context, and model use",
                [
                    ("parameter", "Parameter"),
                    ("estimate", "Estimate"),
                    ("lower", "Lower"),
                    ("upper", "Upper"),
                    ("unit", "Unit"),
                    ("context", "Context"),
                    ("used_as", "Used as"),
                ],
            ),
            (
                "Table S1B. Sources and evidence notes",
                [
                    ("parameter", "Parameter"),
                    ("source", "Source"),
                    ("DOI_or_URL", "DOI or URL"),
                    ("access_date", "Access date"),
                    ("evidence_quality", "Evidence quality"),
                    ("notes", "Notes"),
                ],
            ),
        ],
        SUPPLEMENT_TABLE_DIR / "table_s1_full_parameter_provenance.docx",
    )
    endpoint_rows = [
        ["Distance", "Route definition", "Structural"],
        ["Turning", "Route definition", "Structural"],
        [
            "Movement starts/stops",
            "Progression/control rule plus congestion",
            "Model-dependent dynamic",
        ],
        ["Throughput", "Arrival/service/route/control interaction", "Dynamic"],
        ["Waiting", "Arrival/service/route/control interaction", "Dynamic"],
        [
            "Required waiting area",
            "Demand plus spatial assumptions",
            "Conditional engineering outcome",
        ],
    ]
    endpoint_headers = ["Endpoint", "Mechanism", "Classification"]
    with (SUPPLEMENT_TABLE_DIR / "table_s2_endpoint_classification.csv").open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(endpoint_headers)
        writer.writerows(endpoint_rows)
    _table_document(
        "Table S2. Endpoint mechanism and interpretation",
        endpoint_headers,
        endpoint_rows,
        SUPPLEMENT_TABLE_DIR / "table_s2_endpoint_classification.docx",
    )
    architectures = [
        ["S0", "Conventional switchback", "Continuous/standard release", "B0"],
        ["S1", "Stationary rows", "Full-row batches", "B1–B4"],
        ["S2", "Stationary rows", "Sub-batches", "B1–B4"],
        ["S3", "Stationary rows", "Downstream threshold", "B1–B4"],
        ["S4", "Stationary rows", "Noisy capacity prediction", "B1–B4"],
        ["S5", "Virtual queue benchmark", "Called release", "B3"],
    ]
    with (TABLE_DIR / "table_02_architectures.csv").open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["Policy", "Waiting architecture", "Control rule", "Layouts"])
        writer.writerows(architectures)
    _table_document(
        "Table 2. Architectures and control policies",
        ["Policy", "Waiting architecture", "Control rule", "Layouts"],
        architectures,
        TABLE_DIR / "table_02_architectures.docx",
    )
    primary_path = Path("results/phase5/primary_results.csv")
    primary = pd.read_csv(primary_path)
    main_endpoints = [
        "mean_movement_distance_m",
        "mean_movement_starts",
        "mean_turning_angle_degrees",
        "required_waiting_area_m2_at_density_limit",
        "throughput_per_hour",
    ]
    endpoint_labels = {
        "mean_movement_distance_m": "Movement distance (m/passenger)",
        "mean_movement_starts": "Modeled movement starts (per passenger)",
        "mean_turning_angle_degrees": "Cumulative turning (degrees/passenger)",
        "required_waiting_area_m2_at_density_limit": "Required waiting area (m²)",
        "throughput_per_hour": "Throughput (passengers/h)",
    }
    endpoint_digits = {
        "mean_movement_distance_m": 1,
        "mean_movement_starts": 1,
        "mean_turning_angle_degrees": 0,
        "required_waiting_area_m2_at_density_limit": 1,
        "throughput_per_hour": 1,
    }
    endpoint_interpretations = {
        "mean_movement_distance_m": "Structural route contrast",
        "mean_movement_starts": "Model-dependent; fewer starts",
        "mean_turning_angle_degrees": "Structural route contrast",
        "required_waiting_area_m2_at_density_limit": (
            "Conditional spatial outcome; lower area"
        ),
        "throughput_per_hour": "Safeguard met: lower CI > −2% of S0",
    }
    table_rows: list[list[str]] = []
    for load in sorted(primary["load_regime"].unique()):
        for endpoint in main_endpoints:
            row = primary[
                (primary["load_regime"] == load)
                & (primary["endpoint"] == endpoint)
            ].iloc[0]
            digits = endpoint_digits[endpoint]
            table_rows.append(
                [
                    f"{load:.2f}",
                    endpoint_labels[endpoint],
                    _display_number(float(row["baseline_mean"]), digits),
                    _display_number(
                        float(row["mean_paired_difference"]),
                        digits,
                    ),
                    (
                        f"{_display_number(float(row['ci95_lower']), digits)} to "
                        f"{_display_number(float(row['ci95_upper']), digits)}"
                    ),
                    endpoint_interpretations[endpoint],
                ]
            )
    table_headers = [
        "Load",
        "Endpoint",
        "S0 mean",
        "S3−S0 difference",
        "95% paired CI",
        "Interpretation/criterion",
    ]
    with (TABLE_DIR / "table_03_primary_outcomes.csv").open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(table_headers)
        writer.writerows(table_rows)
    _table_document(
        "Table 3. Primary paired outcomes and interpretation",
        table_headers,
        table_rows,
        TABLE_DIR / "table_03_primary_outcomes.docx",
    )
    supplementary_primary_path = (
        SUPPLEMENT_TABLE_DIR / "table_s3_locked_primary_audit.csv"
    )
    shutil.copyfile(primary_path, supplementary_primary_path)
    with primary_path.open(newline="", encoding="utf-8") as handle:
        audit_rows = list(csv.DictReader(handle))
    _multi_panel_table_document(
        "Table S3. Complete locked-primary audit table",
        audit_rows,
        [
            (
                "Table S3A. Design and paired estimates",
                [
                    ("load_regime", "Load"),
                    ("policy", "Policy"),
                    ("baseline", "Baseline"),
                    ("endpoint", "Endpoint"),
                    ("replications", "N"),
                    ("baseline_mean", "S0 mean"),
                    ("mean_paired_difference", "S3−S0"),
                ],
            ),
            (
                "Table S3B. Paired uncertainty and relative bound",
                [
                    ("load_regime", "Load"),
                    ("endpoint", "Endpoint"),
                    ("standard_deviation_paired_difference", "Paired SD"),
                    ("ci95_lower", "95% CI lower"),
                    ("ci95_upper", "95% CI upper"),
                    ("ci95_half_width", "95% half-width"),
                    ("relative_lower_bound", "Relative lower bound"),
                ],
            ),
            (
                "Table S3C. Throughput safeguard decision fields",
                [
                    ("load_regime", "Load"),
                    ("endpoint", "Endpoint"),
                    ("throughput_acceptable", "Locked decision"),
                    ("throughput_acceptable_margin_0pct", "0% margin"),
                    ("throughput_acceptable_margin_1pct", "1% margin"),
                    ("throughput_acceptable_margin_2pct", "2% margin"),
                    ("throughput_acceptable_margin_5pct", "5% margin"),
                ],
            ),
        ],
        SUPPLEMENT_TABLE_DIR / "table_s3_locked_primary_audit.docx",
    )
    robustness = pd.read_csv("results/phase8/scenario_comparisons.csv")
    counts = Counter()
    for labels in robustness["failure_labels"]:
        for label in str(labels).split(";"):
            if label != "none":
                counts[label] += 1
    robust_rows = [
        [label.replace("_", " "), str(count), str(len(robustness))]
        for label, count in sorted(counts.items())
    ]
    with (TABLE_DIR / "table_04_failure_summary.csv").open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["Criterion", "Scenarios triggering", "Scenarios tested"])
        writer.writerows(robust_rows)
    _table_document(
        "Table 4. Robustness and failure-region summary",
        ["Criterion", "Scenarios triggering", "Scenarios tested"],
        robust_rows,
        TABLE_DIR / "table_04_failure_summary.docx",
    )


def _build_deck(
    main_items: list[tuple[Path, str]],
    supplement_items: list[tuple[Path, str]],
) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    items = [
        *[(f"Figure {number}", path, caption) for number, (path, caption) in enumerate(main_items, 1)],
        *[
            (f"Figure S{number}", path, caption)
            for number, (path, caption) in enumerate(supplement_items, 1)
        ],
    ]
    for label, path, caption in items:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.5))
        title.text_frame.paragraphs[0].text = label
        title.text_frame.paragraphs[0].font.size = Pt(24)
        slide.shapes.add_picture(str(path), Inches(0.6), Inches(0.8), width=Inches(12.1))
        box = slide.shapes.add_textbox(Inches(0.6), Inches(6.8), Inches(12.0), Inches(0.45))
        box.text_frame.paragraphs[0].text = caption
        box.text_frame.paragraphs[0].font.size = Pt(11)
    presentation.save(FIGURE_DIR / "all_figures_editable.pptx")


def main() -> None:
    _prepare()
    movement_main, movement_supplement = _movement_load_figures()
    main_items = [
        _concept_figure(),
        _trajectory_figure(),
        movement_main,
        _control_response_figure(),
        _space_figure(),
        _pareto_figure(),
        _phase_map_figure(),
    ]
    supplement_items = [*_supplementary_figures(), movement_supplement]
    _build_tables()
    _build_deck(main_items, supplement_items)
    caption_lines = ["# Figure captions", ""]
    for number, (_, caption) in enumerate(main_items, start=1):
        caption_lines.extend([f"**Figure {number}.** {caption}", ""])
    for number, (_, caption) in enumerate(supplement_items, start=1):
        caption_lines.extend([f"**Figure S{number}.** {caption}", ""])
    Path("FIGURE_CAPTIONS.md").write_text(
        "\n".join(caption_lines),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
